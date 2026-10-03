#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Core document builder: assembles sections from the codes registry into a
formatted Word document with TOC placeholder, page numbers, and owner style.
"""

import logging
import time
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from engine.dependency_resolver import DependencyResolver
from engine.formatter import Formatter
from engine.logger import generate_audit_trail
from engine.style_applier import StyleApplier
from engine.validator import Validator
from utils.content_library import ContentLibrary
from utils.docx_manipulator import (
    add_trial_watermark,
    enforce_document_rtl,
    new_docx,
    save_docx,
    set_rtl,
)
from utils.image_processor import embed_image
from utils.json_manager import load_json

logger = logging.getLogger(__name__)

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
_METADATA_DIR = _PROJECT_ROOT
_STYLE_DIR = _PROJECT_ROOT / "templates" / "style_templates"
_SOURCE_DOCS_DIR = _PROJECT_ROOT / "templates" / "source_documents"
_MASTER_CONFIG = _PROJECT_ROOT / "master_config.json"
_OUTPUT_FONT_FAMILY = "Tajawal"


class Builder:
    """
    Assembles a technical proposal Word document from selected activity codes.

    Usage:
        registry = load_json("codes_registry.json")
        builder = Builder(registry["codes"])
        builder.build(
            selected_codes=["001-SUR-BASE", "002-EXC-FINE", ...],
            project_id="wastewater",
            owner_id="nwc",
            output_path="output/generated_documents/proposal.docx"
        )
    """

    def __init__(
        self,
        codes: Dict[str, Dict],
        metadata_dir: str | Path = _METADATA_DIR,
        style_dir: str | Path = _STYLE_DIR,
        source_docs_dir: str | Path = _SOURCE_DOCS_DIR,
        master_config: str | Path = _MASTER_CONFIG,
        *,
        validator: Optional[Validator] = None,
        resolver: Optional[DependencyResolver] = None,
        style_applier: Optional[StyleApplier] = None,
        content_lib: Optional[ContentLibrary] = None,
    ):
        """Initialise the builder with a codes registry and optional collaborators.

        Positional / keyword-only directory arguments are used only when the
        corresponding collaborator is *not* supplied.  Inject pre-built instances
        to override behaviour in tests or to share expensive objects.

        Args:
            codes:          Full codes registry dict (``registry["codes"]``).
            metadata_dir:   Directory that holds ``<project_id>_project_metadata.json``.
            style_dir:      Directory that holds ``<owner_id>_style.json`` files.
            source_docs_dir:Directory that holds per-code ``.docx`` source files.
            master_config:  Path to ``master_config.json`` (used to resolve owner names).
            validator:      Optional pre-built :class:`~engine.validator.Validator`.
            resolver:       Optional pre-built :class:`~engine.dependency_resolver.DependencyResolver`.
            style_applier:  Optional pre-built :class:`~engine.style_applier.StyleApplier`.
            content_lib:    Optional pre-built :class:`~utils.content_library.ContentLibrary`.
        """
        self._codes = codes
        self._metadata_dir = Path(metadata_dir)
        self._style_dir = Path(style_dir)
        self._validator = validator if validator is not None else Validator(codes)
        self._resolver = resolver if resolver is not None else DependencyResolver(codes)
        self._style_applier = (
            style_applier if style_applier is not None else StyleApplier(style_dir)
        )
        self._content_lib = (
            content_lib if content_lib is not None else ContentLibrary(source_docs_dir)
        )
        self._project_metadata: Dict[str, Dict] = {}

        # Load owner display names from master_config.json (owner_specifications section)
        try:
            cfg = load_json(Path(master_config))
            self._owner_specs: Dict[str, Dict] = cfg.get("owner_specifications", {})
            # Category ranges (001–005) used for section grouping + dividers
            self._code_ranges: Dict[str, Dict] = cfg.get("code_ranges", {})
        except (FileNotFoundError, ValueError, OSError):
            logger.warning("master_config.json not loadable — owner names will fall back to owner_id")
            self._owner_specs = {}
            self._code_ranges = {}

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def build(
        self,
        selected_codes: List[str],
        project_id: str,
        owner_id: str,
        output_path: str | Path,
        skip_validation: bool = False,
        boq_order: Optional[List[str]] = None,
        template_vars: Optional[Dict[str, str]] = None,
    ) -> Tuple[bool, Optional[str]]:
        """
        Build the proposal document.

        Returns:
            (success, error_message_ar)
            error_message_ar is None on success.

        Validation warnings (e.g. missing mandatory codes) are logged but do
        not abort the build.  Hard errors (forbidden codes, inactive codes)
        cause build failure.
        """
        output_path = Path(output_path)
        start = time.monotonic()

        # Defensive normalization: accept str or single-item list (Validator
        # accepts both); a filename must never be derived from a list.
        if not isinstance(project_id, str):
            project_id = str(project_id[0]) if project_id else ""

        # --- Validate ---
        if not skip_validation:
            is_valid, errors, warnings = self._validator.validate(
                selected_codes, owner_id, project_id
            )
            for w in warnings:
                logger.warning("Validation warning: %s", w)
            if not is_valid:
                msg = "فشل التحقق من الأكواد:\n" + "\n".join(f"• {e}" for e in errors)
                logger.error("Build aborted — validation failed:\n%s", "\n".join(errors))
                self._write_audit(
                    selected_codes, project_id, owner_id, output_path,
                    "failure", msg, 0, 0
                )
                return False, msg

        # --- Resolve & order codes ---
        if boq_order:
            ordered_codes = self._resolver.resolve_with_order(selected_codes, boq_order)
        else:
            ordered_codes = self._resolver.resolve(selected_codes)
        project_meta = self._load_project_metadata(project_id)

        # --- Assemble document ---
        doc = new_docx()
        style_spec = self._load_style_spec(owner_id)
        formatter = Formatter(style_spec) if style_spec else None

        self._add_cover(doc, project_id, owner_id, ordered_codes, formatter, template_vars)
        self._add_toc_placeholder(doc)
        self._add_sections(doc, ordered_codes, project_meta, formatter)
        self._style_applier.apply_style(doc, owner_id)

        # --- علامة مائية للنسخة التجريبية (ترخيص يوم واحد) ---
        try:
            from utils.license_manager import check_saved_license
            lic = check_saved_license()
            if lic.get("days_left") is not None and lic["days_left"] <= 1:
                add_trial_watermark(doc)
                logger.info("Trial watermark added (days_left=%s)", lic["days_left"])
        except Exception:
            pass   # لا يوقف البناء إذا فشل فحص الترخيص

        self._finalize_arabic_layout(doc)

        # --- Save ---
        save_docx(doc, output_path)

        elapsed = time.monotonic() - start
        file_size = output_path.stat().st_size if output_path.exists() else 0

        logger.info(
            "Build complete: %s  (%.2fs, %d bytes)",
            output_path, elapsed, file_size
        )
        self._write_audit(
            selected_codes, project_id, owner_id, output_path,
            "success", None, elapsed, file_size
        )
        return True, None

    @staticmethod
    def _finalize_arabic_layout(doc: Document) -> None:
        """Apply the final non-negotiable Arabic output rules before saving."""
        from utils.docx_manipulator import (
            apply_font_family_to_document,
            set_document_default_font,
        )

        Builder._apply_doc_rtl(doc)
        apply_font_family_to_document(doc, _OUTPUT_FONT_FAMILY)
        set_document_default_font(doc, _OUTPUT_FONT_FAMILY)
        logger.debug("Applied output font '%s' to the full document", _OUTPUT_FONT_FAMILY)

    def estimate_pages(self, selected_codes: List[str]) -> int:
        """Estimate total page count for the selected codes.

        Includes 2 fixed pages present in every proposal:
          +1 cover page
          +1 table of contents

        Plus +1 divider page for every extra category represented
        (the first category's divider shares the previous page boundary
        — only transitions add a full page).
        """
        code_pages = sum(
            self._codes[c].get("page_count", 0)
            for c in selected_codes
            if c in self._codes
        )
        divider_pages = self._category_divider_count(selected_codes)
        return code_pages + 2 + divider_pages

    def estimate_images(self, selected_codes: List[str]) -> int:
        """Estimate total image count for the selected codes."""
        return sum(
            self._codes[c].get("image_count", 0)
            for c in selected_codes
            if c in self._codes
        )

    def content_summary(self, selected_codes: List[str]) -> Dict[str, str]:
        """
        Return a dict mapping each code to its content status:
            "library"     → real .docx file found in content library
            "placeholder" → no source file yet (will use metadata text)
        """
        return {
            code_id: ("library" if self._content_lib.exists(code_id) else "placeholder")
            for code_id in selected_codes
            if code_id in self._codes
        }

    # ------------------------------------------------------------------
    # Document assembly
    # ------------------------------------------------------------------

    def _add_cover(
        self,
        doc: Document,
        project_id: str,
        owner_id: str,
        ordered_codes: List[str],
        formatter: Optional[Formatter],
        template_vars: Optional[Dict[str, str]] = None,
    ) -> None:
        """Render a centred cover page: title, project, owner, date, and stats.

        If *template_vars* is provided, any non-empty value in it overrides
        the corresponding metadata / default:

        - ``project_name``    — overrides the metadata project name
        - ``tender_number``   — shown as "رقم المنافسة: <value>"
        - ``submission_date`` — overrides today's date
        - ``contract_value``  — shown as "قيمة العقد: <value>"
        - ``engineer_name``   — shown as "إعداد: <value>"
        """
        from datetime import date
        tv = template_vars or {}

        project_meta = self._load_project_metadata(project_id)
        project_name = tv.get("project_name") or project_meta.get("project_metadata", {}).get(
            "name_ar", project_id
        )

        owner_name = self._owner_name_ar(owner_id)

        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_para.add_run("عرض فني")
        run.font.size = Pt(24)
        run.bold = True

        subtitle_para = doc.add_paragraph()
        subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_para.add_run(project_name).font.size = Pt(16)

        owner_para = doc.add_paragraph()
        owner_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        owner_para.add_run(f"مُقدَّم إلى: {owner_name}").font.size = Pt(14)

        # Tender number — shown only when provided
        tender_number = tv.get("tender_number", "")
        if tender_number:
            tender_para = doc.add_paragraph()
            tender_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tender_para.add_run(f"رقم المنافسة: {tender_number}").font.size = Pt(12)

        # Submission date — custom value or today
        submission_date = tv.get("submission_date") or date.today().strftime("%Y-%m-%d")
        date_para = doc.add_paragraph()
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_para.add_run(f"التاريخ: {submission_date}").font.size = Pt(12)

        # Contract value — shown only when provided
        contract_value = tv.get("contract_value", "")
        if contract_value:
            contract_para = doc.add_paragraph()
            contract_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contract_para.add_run(f"قيمة العقد: {contract_value}").font.size = Pt(12)

        # Engineer name — shown only when provided
        engineer_name = tv.get("engineer_name", "")
        if engineer_name:
            engineer_para = doc.add_paragraph()
            engineer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            engineer_para.add_run(f"إعداد: {engineer_name}").font.size = Pt(12)

        pages = self.estimate_pages(ordered_codes)
        stats_para = doc.add_paragraph()
        stats_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        stats_para.add_run(
            f"عدد الأكواد: {len(ordered_codes)}  |  الصفحات التقديرية: {pages}"
        ).font.size = Pt(11)

        doc.add_page_break()

    def _add_toc_placeholder(self, doc: Document) -> None:
        """Insert a TOC field that Word regenerates on open."""
        toc_heading = doc.add_heading("فهرس المحتويات", level=1)
        toc_heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Insert TOC field via raw OOXML
        para = doc.add_paragraph()
        run = para.add_run()
        fldChar_begin = OxmlElement("w:fldChar")
        fldChar_begin.set(qn("w:fldCharType"), "begin")
        fldChar_begin.set(qn("w:dirty"), "true")

        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = ' TOC \\o "1-3" \\h \\z \\u '

        fldChar_end = OxmlElement("w:fldChar")
        fldChar_end.set(qn("w:fldCharType"), "end")

        run._r.append(fldChar_begin)
        run._r.append(instrText)
        run._r.append(fldChar_end)

        doc.add_page_break()

    def _add_sections(
        self,
        doc: Document,
        ordered_codes: List[str],
        project_meta: Dict,
        formatter: Optional[Formatter],
    ) -> None:
        """
        Add one section per code in sequence order, grouped by category.

        A category divider (title + description on its own page) is inserted
        before the first code of each new category (001–005), turning the flat
        code stack into ordered proposal phases.

        Priority per section:
          1. Real content from content library (templates/source_documents/{code_id}.docx)
          2. Metadata-based placeholder (activity name + description)
        """
        activities = {
            a["code_id"]: a
            for a in project_meta.get("activity_sequence", [])
            if isinstance(a, dict) and "code_id" in a
        }

        current_category: Optional[str] = None
        for code_id in ordered_codes:
            code = self._codes.get(code_id)
            if not code:
                continue

            category = code.get("category")
            if category != current_category:
                self._add_category_divider(
                    doc, category, formatter,
                    is_first=(current_category is None),
                )
                current_category = category

            activity = activities.get(code_id, {})
            name_ar = code.get("activity_name_ar", code_id)
            name_en = code.get("activity_name_en", "")

            # Section heading — always shown regardless of content source
            # Arabic name leads (RTL convention): name_ar — code_id
            heading_text = f"{name_ar}  —  {code_id}"
            if formatter:
                formatter.add_heading(doc, heading_text, level=1)
            else:
                h = doc.add_heading(heading_text, level=1)
                h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(h)

            # Try content library first
            has_real_content = self._content_lib.insert_into(doc, code_id)

            if not has_real_content:
                # Fallback: placeholder from metadata
                if name_en:
                    if formatter:
                        formatter.add_heading(doc, name_en, level=3)
                    else:
                        doc.add_heading(name_en, level=3)

                description = activity.get("description", "")
                if description:
                    if formatter:
                        formatter.add_body_paragraph(doc, description)
                    else:
                        doc.add_paragraph(description)
            else:
                logger.info("Used content library for %s", code_id)

            doc.add_paragraph()  # section spacer

    # ------------------------------------------------------------------
    # Category grouping helpers
    # ------------------------------------------------------------------

    def _category_divider_count(self, selected_codes: List[str]) -> int:
        """Number of full divider pages: every category transition after the first."""
        cats: List[str] = []
        for cid in selected_codes:
            code = self._codes.get(cid)
            if not code:
                continue
            cat = code.get("category")
            if cat != (cats[-1] if cats else None):
                cats.append(cat)
        return max(0, len(cats) - 1)

    def _category_info(self, category: Optional[str]) -> Dict:
        """Resolve a category id to its display info (name + description).

        Falls back to the raw id / empty text when the category is unknown
        (e.g. custom 999-CUS codes) so existing builds never break.
        """
        info = self._code_ranges.get(category, {}) if category else {}
        return {
            "name_ar": info.get("category_name_ar") or category or "أخرى",
            "description": info.get("description", ""),
        }

    def _add_category_divider(
        self,
        doc: Document,
        category: Optional[str],
        formatter: Optional[Formatter],
        *,
        is_first: bool = False,
    ) -> None:
        """Insert a category divider: page break + title + description.

        The first category does NOT get a leading page break (it follows the
        TOC page naturally); every transition after it starts a new page.
        """
        if not is_first:
            doc.add_page_break()

        info = self._category_info(category)
        name_ar = info["name_ar"]
        description = info["description"]

        if formatter:
            formatter.add_heading(doc, name_ar, level=1)
            if description:
                formatter.add_body_paragraph(doc, description)
        else:
            h = doc.add_heading(name_ar, level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(h)
            if description:
                p = doc.add_paragraph(description)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                set_rtl(p)

        doc.add_paragraph()  # spacer after divider
        logger.info("Category divider added: %s", name_ar)

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def _owner_name_ar(self, owner_id: str) -> str:
        """Return the Arabic display name for *owner_id* from master_config.json.

        Falls back to ``owner_id`` itself when the config is absent or the
        owner is not yet registered, so existing builds never break.
        """
        spec = self._owner_specs.get(owner_id, {})
        return spec.get("owner_name_ar", owner_id)

    def _load_project_metadata(self, project_id: str) -> Dict:
        """Load ``<project_id>_project_metadata.json``; returns empty template on miss."""
        if isinstance(project_id, (list, tuple)):
            project_id = str(project_id[0]) if project_id else ""
        if project_id in self._project_metadata:
            return self._project_metadata[project_id]
        path = self._metadata_dir / f"{project_id}_project_metadata.json"
        try:
            meta = load_json(path)
        except FileNotFoundError:
            meta = {"activity_sequence": []}
        self._project_metadata[project_id] = meta
        return meta

    def _load_style_spec(self, owner_id: str) -> Optional[Dict]:
        """Load ``<owner_id>_style.json`` from style_dir; returns None if absent."""
        path = self._style_dir / f"{owner_id}_style.json"
        try:
            return load_json(path)
        except FileNotFoundError:
            return None

    @staticmethod
    def _apply_doc_rtl(doc: Document) -> None:
        """Enforce RTL direction after full assembly.

        Called as the last step before saving — catches:
        - Paragraphs from content-library ``.docx`` files (may have LTR source).
        - Fallback paragraphs added directly via ``doc.add_paragraph()`` when
          no :class:`~engine.formatter.Formatter` is available.
        - Any paragraphs added by :class:`~engine.style_applier.StyleApplier`.
        - Table-cell paragraphs, headers, and footers.

        CENTER-aligned paragraphs (cover page title/stats) are preserved as-is.
        """
        enforce_document_rtl(doc)

    def _write_audit(
        self,
        codes: List[str],
        project_id: str,
        owner_id: str,
        output_path: Path,
        status: str,
        error: Optional[str],
        elapsed: float,
        file_size: int,
    ) -> None:
        """Persist an audit-trail entry via engine.logger.generate_audit_trail.

        Failures are logged as warnings rather than raised — a broken audit
        trail must never abort a successful build.
        """
        try:
            generate_audit_trail({
                "selected_codes": codes,
                "project_id": project_id,
                "owner_id": owner_id,
                "output_path": str(output_path),
                "status": status,
                "error": error,
                "processing_time_seconds": round(elapsed, 3),
                "file_size_bytes": file_size,
            })
        except Exception as exc:
            logger.warning("Audit trail write failed: %s", exc)
