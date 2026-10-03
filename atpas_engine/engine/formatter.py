#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Applies paragraph styles, RTL direction, headings, and spacing to a Document.
Operates on a docx.Document object; does not read/write files.
"""

from typing import Dict, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from utils.docx_manipulator import set_paragraph_font, set_rtl, _hex_to_rgb


class Formatter:
    """
    Apply typographic styles to a Document based on a style spec dict.

    The style spec mirrors the JSON structure in templates/style_templates/*.json:
        {
          "fonts": {"body": {"family": ..., "size": ..., "bold": ...}, "heading1": {...}, ...},
          "colors": {"heading": "#RRGGBB", "text": "#RRGGBB"},
          "rtl_direction": true
        }
    """

    def __init__(self, style_spec: Dict):
        self._spec = style_spec
        self._fonts = style_spec.get("fonts", {})
        self._colors = style_spec.get("colors", {})
        self._rtl = style_spec.get("rtl_direction", True)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def format_document(self, doc: Document) -> None:
        """Apply all formatting rules to every paragraph in the document."""
        for para in doc.paragraphs:
            self._format_paragraph(para)

    def format_paragraph(self, para: Paragraph, style_key: str = "body") -> None:
        """Format a single paragraph using the named style key."""
        self._apply_font(para, style_key)
        if self._rtl:
            set_rtl(para)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def add_heading(
        self,
        doc: Document,
        text: str,
        level: int = 1,
        color_hex: Optional[str] = None,
    ) -> Paragraph:
        """Add a styled heading paragraph to the document."""
        para = doc.add_heading(text, level=level)
        style_key = f"heading{level}"
        self._apply_font(para, style_key, color_override=color_hex)
        if self._rtl:
            set_rtl(para)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return para

    def add_body_paragraph(self, doc: Document, text: str) -> Paragraph:
        """Add a styled body paragraph."""
        para = doc.add_paragraph(text)
        self._apply_font(para, "body")
        if self._rtl:
            set_rtl(para)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return para

    def set_paragraph_spacing(
        self,
        para: Paragraph,
        space_before_pt: float = 6,
        space_after_pt: float = 6,
        line_spacing_pt: Optional[float] = None,
    ) -> None:
        pf = para.paragraph_format
        pf.space_before = Pt(space_before_pt)
        pf.space_after = Pt(space_after_pt)
        if line_spacing_pt:
            pf.line_spacing = Pt(line_spacing_pt)

    # ------------------------------------------------------------------
    # Internal
    # ------------------------------------------------------------------

    def _format_paragraph(self, para: Paragraph) -> None:
        style_name = para.style.name if para.style else "Normal"
        if style_name.startswith("Heading"):
            try:
                level = int(style_name.split()[-1])
            except (ValueError, IndexError):
                level = 1
            style_key = f"heading{min(level, 3)}"
            color = self._colors.get("heading")
        else:
            style_key = "body"
            color = self._colors.get("text")

        self._apply_font(para, style_key, color_override=color)
        if self._rtl:
            set_rtl(para)
            # Preserve CENTER alignment (cover page title / stats paragraphs)
            if para.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def _apply_font(
        self,
        para: Paragraph,
        style_key: str,
        color_override: Optional[str] = None,
    ) -> None:
        spec = self._fonts.get(style_key, self._fonts.get("body", {}))
        family = spec.get("family", "Arial")
        size = spec.get("size", 12)
        bold = spec.get("bold", False)
        italic = spec.get("italic", False)
        color = color_override or self._colors.get("text")
        set_paragraph_font(para, family, size, bold=bold, italic=italic, color_hex=color)
