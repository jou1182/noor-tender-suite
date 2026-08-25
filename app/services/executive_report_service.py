import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

class ExecutiveReportService:
    @staticmethod
    def generate_pdf_report(tender_data: dict, output_path: str):
        doc = SimpleDocTemplate(output_path, pagesize=letter)
        styles = getSampleStyleSheet()
        Story = []
        
        Story.append(Paragraph(f"Executive Audit Report: Tender #{tender_data.get('id', 'N/A')}", styles["Title"]))
        Story.append(Spacer(1, 12))
        Story.append(Paragraph(f"Score: {tender_data.get('score', 0)}", styles["Heading2"]))
        Story.append(Spacer(1, 12))
        
        Story.append(Paragraph("Compliance Summary", styles["Heading2"]))
        for rec in tender_data.get('records', []):
            text = f"<b>{rec['clause_code']}:</b> {rec['status']} - {rec.get('gap_analysis', '')}"
            Story.append(Paragraph(text, styles["Normal"]))
            Story.append(Spacer(1, 6))
            
        doc.build(Story)

    @staticmethod
    def generate_rfi_docx(rfis: list, output_path: str):
        doc = Document()
        doc.add_heading('Request For Information (RFI) Register', 0)
        
        for idx, rfi in enumerate(rfis, 1):
            doc.add_heading(f'RFI #{idx}', level=1)
            doc.add_paragraph(rfi)
            
        doc.save(output_path)
